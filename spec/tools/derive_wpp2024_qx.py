#!/usr/bin/env python3
"""Derive q_x from UN WPP 2024 exact-age life-table survivor workbooks.

Download the three official MORT/F04 files listed in data/mortality-estimates.json
into --input-dir. This script writes the five configured country series and
calendar years 1950-2100. Annual q_x through age 99 comes from WPP l(x). The
open age interval 100+ is a low-confidence game estimate extrapolated from the
logit trend in q_98 and q_99; it is not a WPP value.

Requires openpyxl. Source release attribution and license are recorded in the
data manifest and source workbook notes.
"""

import argparse
import csv
import gzip
from pathlib import Path

from openpyxl import load_workbook


SOURCES = {
    "both": "WPP2024_MORT_F04_1_LIFE_TABLE_SURVIVORS_BOTH_SEXES.xlsx",
    "male": "WPP2024_MORT_F04_2_LIFE_TABLE_SURVIVORS_MALE.xlsx",
    "female": "WPP2024_MORT_F04_3_LIFE_TABLE_SURVIVORS_FEMALE.xlsx",
}
LOCATIONS = {
    "Japan": "jp_modern",
    "China": "cn_modern",
    "India": "in_modern",
    "Mali": "ml_modern",
    "France": "fr_modern",
}
FIELDS = [
    "scenario_id",
    "country_name_wpp",
    "sex_series",
    "calendar_year",
    "series_type",
    "age_start_years",
    "age_interval_years",
    "q_x",
    "source_measure",
    "source_version",
]


def selected_rows(worksheet, first_year, last_year):
    worksheet.reset_dimensions()
    for row in worksheet.iter_rows(values_only=True):
        if len(row) < 112:
            continue
        country, year = row[2], row[10]
        if country not in LOCATIONS or not isinstance(year, (int, float)):
            continue
        if first_year <= year <= last_year:
            yield country, int(year), row[11:112]


def logit(probability):
    import math
    epsilon = 1e-12
    probability = min(1 - epsilon, max(epsilon, probability))
    return math.log(probability / (1 - probability))


def logistic(value):
    import math
    return 1 / (1 + math.exp(-value))


def derive(source: Path, output: Path, sex: str) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    count = 0
    try:
        with gzip.open(output, "wt", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=FIELDS)
            writer.writeheader()
            ranges = [
                ("Estimates", "estimate", 1950, 2023),
                ("Medium variant", "medium_projection", 2024, 2100),
            ]
            for sheet_name, series_type, first, last in ranges:
                for country, year, survivors in selected_rows(
                    workbook[sheet_name], first, last
                ):
                    if len(survivors) != 101 or any(v is None for v in survivors):
                        raise ValueError(f"Incomplete l(x) row: {country}, {year}")
                    for age in range(100):
                        lx, lx_next = float(survivors[age]), float(survivors[age + 1])
                        if lx <= 0:
                            raise ValueError(f"Invalid l(x): {country}, {year}, {age}")
                        qx = 1.0 - lx_next / lx
                        if not 0 <= qx <= 1:
                            raise ValueError(f"Invalid l(x): {country}, {year}, {age}")
                        writer.writerow(
                            {
                                "scenario_id": LOCATIONS[country],
                                "country_name_wpp": country,
                                "sex_series": sex,
                                "calendar_year": year,
                                "series_type": series_type,
                                "age_start_years": age,
                                "age_interval_years": 1,
                                "q_x": format(qx, ".10g"),
                                "source_measure": "1 - l(x+1)/l(x)",
                                "source_version": "WPP 2024",
                            }
                        )
                        count += 1
                    q98 = 1.0 - float(survivors[99]) / float(survivors[98])
                    q99 = 1.0 - float(survivors[100]) / float(survivors[99])
                    trend = max(0.0, logit(q99) - logit(q98))
                    q100_plus = min(0.99, logistic(logit(q99) + trend))
                    writer.writerow(
                        {
                            "scenario_id": LOCATIONS[country],
                            "country_name_wpp": country,
                            "sex_series": sex,
                            "calendar_year": year,
                            "series_type": series_type,
                            "age_start_years": 100,
                            "age_interval_years": "open",
                            "q_x": format(q100_plus, ".10g"),
                            "source_measure": "estimated open age tail from logit trend in q98-q99",
                            "source_version": "WPP 2024 plus game estimate",
                        }
                    )
                    count += 1
    finally:
        workbook.close()
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("data"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    for sex, filename in SOURCES.items():
        source = args.input_dir / filename
        if not source.is_file():
            raise FileNotFoundError(source)
        output = args.output_dir / f"mortality-wpp2024-qx-{sex}.csv.gz"
        rows = derive(source, output, sex)
        print(f"{output}: {rows} age-year records")


if __name__ == "__main__":
    main()
